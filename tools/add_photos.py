# -*- coding: utf-8 -*-
"""
Проставляет в прайс Яндекса ссылки на кадры программ.

Яндекс забирает картинку по ссылке, поэтому кадры выложены на сайт в JPEG
(папка img-jpg) — WebP он может не понять. У позиций, где фото уже есть
(их Яндекс хранит у себя), ссылка не трогается.
"""
import os, re, json
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, '..', 'price-final.xlsx')
OUT = os.path.join(HERE, '..', 'price-photos.xlsx')
DATA = os.path.join(HERE, '..', 'js', 'data.js')
BASE = 'https://loko-tracker.github.io/7d-attraction/img-jpg/'

src = open(DATA, encoding='utf-8').read()
img_of = {}
for m in re.finditer(r'\{\s*t:\s*"([^"]+)"[^}]*?img:\s*"([^"]+)"', src):
    img_of[m.group(1)] = m.group(2)

wb = openpyxl.load_workbook(SRC)
ws = wb.active
col = {c.value: i + 1 for i, c in enumerate(ws[1])}
C_NAME, C_PHOTO = col['название'], col['фото']

filled, kept, unmatched = [], [], []
for r in range(2, ws.max_row + 1):
    name = ws.cell(r, C_NAME).value
    if ws.cell(r, C_PHOTO).value:
        kept.append(name)
    elif name in img_of:
        ws.cell(r, C_PHOTO, BASE + img_of[name] + '.jpg')
        filled.append(name)
    else:
        unmatched.append(name)

wb.save(os.path.abspath(OUT))
print(json.dumps({
    'строк': ws.max_row - 1,
    'ссылок_проставлено': len(filled),
    'фото_уже_было': len(kept),
    'без_картинки_осталось': unmatched,
    'файл': os.path.abspath(OUT),
}, ensure_ascii=False, indent=1))
