# -*- coding: utf-8 -*-
"""
Проставляет описания в прайс Яндекса и убирает найденного двойника.
Исходник — price-53.xlsx (то, что сейчас залито), результат — price-final.xlsx.
"""
import os, json, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, '..', 'price-53.xlsx')
OUT = os.path.join(HERE, '..', 'price-final.xlsx')
DESC = os.path.join(HERE, 'descriptions.json')

DROP = {'Sky Dragon'}          # дубликат «Парящих храмов»

desc = {k: v for k, v in json.load(open(DESC, encoding='utf-8')).items()
        if not k.startswith('_')}

wb = openpyxl.load_workbook(SRC)
ws = wb.active
col = {c.value: i + 1 for i, c in enumerate(ws[1])}
C_NAME, C_DESC = col['название'], col['описание']

# Удаляем строки-дубликаты (снизу вверх, чтобы не сбить нумерацию).
dropped = []
for r in range(ws.max_row, 1, -1):
    if ws.cell(r, C_NAME).value in DROP:
        dropped.append(ws.cell(r, C_NAME).value)
        ws.delete_rows(r)

filled, already, missing = [], [], []
for r in range(2, ws.max_row + 1):
    name = ws.cell(r, C_NAME).value
    if ws.cell(r, C_DESC).value:
        already.append(name)
    elif name in desc:
        ws.cell(r, C_DESC, desc[name])
        filled.append(name)
    else:
        missing.append(name)

unused = sorted(set(desc) - set(filled))
wb.save(os.path.abspath(OUT))

print(json.dumps({
    'строк_итого': ws.max_row - 1,
    'удалено_дубликатов': dropped,
    'описаний_проставлено': len(filled),
    'уже_было_описаний': len(already),
    'осталось_без_описания': missing,
    'описания_не_пригодились': unused,
    'файл': os.path.abspath(OUT),
}, ensure_ascii=False, indent=1))
